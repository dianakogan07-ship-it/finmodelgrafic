"""Чтение рядов цены/продаж и кэш-значений ФР/LLCR из xlsx-модели по cell_map."""
from openpyxl import load_workbook

from .cellmap import cell_map, col_idx, col_letter


class ModelError(Exception):
    """Модель не прошла валидацию; текст — для админа, с адресом ячейки."""


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def parse(path) -> dict:
    cm = cell_map()
    s = cm["sales"]
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as e:  # битый zip, не xlsx, зашифрованный файл
        raise ModelError(f"Файл не читается как книга Excel ({e.__class__.__name__}).") from e
    try:
        return _parse(wb, cm, s)
    finally:
        wb.close()


def _parse(wb, cm, s):
    warnings = []
    need = {s["sheet"]} | {o["sheet"] for o in cm["outputs"].values()}
    missing = [n for n in need if n not in wb.sheetnames]
    if missing:
        raise ModelError(f"В файле нет листа: {', '.join(missing)}.")

    # --- ФР и LLCR (кэш Excel) ---
    outputs = {}
    for key, o in cm["outputs"].items():
        ws = wb[o["sheet"]]
        v = _num(ws[o["cell"]].value)
        if v is None:
            raise ModelError(f"{o['sheet']}!{o['cell']} пустая или не число ({key}).")
        lbl = ws[o["label_cell"]].value if o.get("label_cell") else None
        if o.get("label") and str(lbl or "").strip() != o["label"]:
            warnings.append(f"{o['sheet']}!{o['label_cell']}: ожидалась подпись «{o['label']}», в файле «{lbl}».")
        outputs[key] = v

    # --- лист продаж ---
    ws = wb[s["sheet"]]
    c_first = col_idx(s["first_col"])
    c_q, c_p, c_t, c_tot = (col_idx(s[k]) for k in ("queue_col", "param_col", "type_col", "total_col"))
    rows = list(ws.iter_rows(min_row=1, max_row=s["scan_rows"], max_col=c_first + 400, values_only=True))

    labels, cols = [], []
    prow = rows[s["periods_row"] - 1] if len(rows) >= s["periods_row"] else ()
    for c in range(c_first, len(prow) + 1):
        v = prow[c - 1]
        if v is None or str(v).strip() == "":
            break
        labels.append(str(v).strip())
        cols.append(col_letter(c))
    if not labels:
        raise ModelError(f"{s['sheet']}!{s['first_col']}{s['periods_row']}: нет подписей кварталов.")

    found: dict[int, dict] = {}
    for r, row in enumerate(rows, start=1):
        cell = lambda c: row[c - 1] if c - 1 < len(row) else None  # noqa: E731
        q, p, t = cell(c_q), str(cell(c_p) or "").strip(), str(cell(c_t) or "").strip()
        if not isinstance(q, (int, float)) or q < 1 or t != s["type"]:
            continue
        q = int(q)
        if p == s["price_param"]:
            found.setdefault(q, {})["price"] = r
        elif p == s["area_param"]:
            found.setdefault(q, {})["area"] = r

    queues = {}
    for q, rr in sorted(found.items()):
        if "price" not in rr or "area" not in rr:
            warnings.append(f"Очередь {q}: найдена только строка {'цены' if 'price' in rr else 'продаж'}.")
            continue
        exp = (s.get("expected_rows") or {}).get(q)
        if exp and (exp["price"], exp["area"]) != (rr["price"], rr["area"]):
            warnings.append(
                f"Очередь {q}: строки найдены в {rr['price']}/{rr['area']}, ожидались {exp['price']}/{exp['area']} — проверьте структуру листа."
            )
        prow_, arow = rows[rr["price"] - 1], rows[rr["area"] - 1]
        get = lambda row, i: _num(row[c_first - 1 + i]) if c_first - 1 + i < len(row) else None  # noqa: E731
        price = [(get(prow_, i) or 0.0) / s["price_divisor"] for i in range(len(labels))]
        area = [get(arow, i) or 0.0 for i in range(len(labels))]
        if not any(a > 0 for a in area):
            continue
        total = sum(area)
        model_total = _num(arow[c_tot - 1]) if c_tot - 1 < len(arow) else None
        if model_total is not None and abs(model_total - total) > 0.01:
            warnings.append(
                f"Очередь {q}: сумма продаж {total:,.2f} м² не сходится с {s['sheet']}!{s['total_col']}{rr['area']} = {model_total:,.2f}.".replace(",", " ")
            )
        queues[str(q)] = {"price": price, "area": area, "total": round(total, 6), "rows": rr}

    if not queues:
        raise ModelError(f"На листе «{s['sheet']}» не найдены строки «{s['price_param']}» / «{s['area_param']}» для «{s['type']}».")

    # окно — от первого до последнего квартала продаж
    active = [i for i in range(len(labels)) if any(qd["area"][i] > 0 for qd in queues.values())]
    a, b = active[0], active[-1] + 1
    for qd in queues.values():
        qd["price"] = [round(x, 6) for x in qd["price"][a:b]]
        qd["area"] = [round(x, 6) for x in qd["area"][a:b]]

    return {
        "labels": labels[a:b],
        "cols": cols[a:b],
        "queues": queues,
        "fr": outputs["fr"],
        "llcr": outputs["llcr"],
        "sheet": s["sheet"],
        "warnings": warnings,
    }
